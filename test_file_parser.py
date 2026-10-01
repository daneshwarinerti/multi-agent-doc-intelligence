"""
test_file_parser.py
-------------------
Verification script for Phase 2, Stage 2 (file_parser.py).
Generates test PDF, DOCX, and TXT files and parses them into page lists.
"""

import os
import pypdf
import docx
from file_parser import parse_file

def create_sample_files():
    os.makedirs("./scratch_test", exist_ok=True)
    
    # 1. Create a sample TXT file
    txt_path = "./scratch_test/sample.txt"
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write("Section 1: Executive Summary\nOur company achieved strong 2023 revenue growth.\n\n"
                "Section 2: Financial Details\nNet margin expanded by 4.2% across key operational regions.")

    # 2. Create a sample DOCX file
    docx_path = "./scratch_test/sample.docx"
    doc = docx.Document()
    doc.add_heading("Q3 Operations Report", level=1)
    doc.add_paragraph("Supply chain efficiency increased by 15% following automated routing upgrades.")
    doc.add_paragraph("Risk mitigation strategies were successfully executed across all regional nodes.")
    doc.save(docx_path)

    # 3. Create a sample PDF file using pypdf writer
    pdf_path = "./scratch_test/sample.pdf"
    writer = pypdf.PdfWriter()
    page1 = writer.add_blank_page(width=612, height=792)
    # Write page 1 using simple annotations or text streams if needed, or parse created PDF
    # For a real PDF test with text, we can use simple blank page + text writing or parse an existing PDF.
    # Note: pypdf doesn't easily create formatted text from scratch without canvas, so we can use TXT & DOCX, 
    # and verify parse_file routing for PDF by writing a minimal text PDF or inspecting writer.
    with open(pdf_path, "wb") as f:
        writer.write(f)

    return txt_path, docx_path, pdf_path

def run_test():
    txt_path, docx_path, pdf_path = create_sample_files()

    print("--- [TEST 1: PARSING PLAIN TXT FILE] ---")
    txt_pages = parse_file(txt_path)
    print(f"File: {txt_path} -> Parsed Pages Count: {len(txt_pages)}")
    print(f"Page 1 Metadata: Page Num = {txt_pages[0]['page_num']}")
    print(f"Page 1 Text Sample:\n\"{txt_pages[0]['text'][:120]}...\"\n")

    print("--- [TEST 2: PARSING DOCX FILE] ---")
    docx_pages = parse_file(docx_path)
    print(f"File: {docx_path} -> Parsed Pages Count: {len(docx_pages)}")
    print(f"Page 1 Metadata: Page Num = {docx_pages[0]['page_num']}")
    print(f"Page 1 Text Sample:\n\"{docx_pages[0]['text'][:120]}...\"\n")

    print("--- [TEST 3: SCANNED / EMPTY PDF FAILURE CASE HANDLER] ---")
    try:
        parse_file(pdf_path)
    except ValueError as e:
        print(f"EXPECTED ERROR CAUGHT: {e}\n")

if __name__ == "__main__":
    run_test()
